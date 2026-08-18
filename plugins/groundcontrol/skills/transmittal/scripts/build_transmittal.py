#!/usr/bin/env python3
"""
build_transmittal.py — render a spaceagency drawing issue register / transmittal (xlsx)
from scan_issues.py JSON. Reproduces the spA template layout (black & white,
KMR Waldenburg Buch Halbschmal). Requires openpyxl.

Usage:
  python3 build_transmittal.py --data data.json --out transmittal.xlsx
     [--project-title "..."] [--project-no NNNN]
"""
import json, re, argparse
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.utils import get_column_letter

FONT = "KMR Waldenburg Buch Halbschmal"

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--project-title", default=None)
    ap.add_argument("--project-no", default=None)
    a = ap.parse_args()
    D = json.load(open(a.data))
    title = a.project_title or D.get("project_title", "")
    pno   = a.project_no or D.get("project_no", "")
    cur, issues = D["current"], D["issues"]

    # exclude DELETED sheets; order by series then number
    def numkey(no):
        m = re.match(r'^[A-Za-z](\d+(?:\.\d+)?)', no); return float(m.group(1)) if m else 9999
    rows = [n for n in cur if "DELETED" not in cur[n]["name"].upper()]
    rows.sort(key=lambda n: (numkey(n), n))
    from collections import OrderedDict
    groups = OrderedDict()
    for n in rows: groups.setdefault(cur[n]["series"], []).append(n)
    groups = OrderedDict(sorted(groups.items(), key=lambda kv: kv[0]))
    issue_only = sorted(set().union(*[set(i["drawings"]) for i in issues]) - set(cur), key=numkey) if issues else []

    wb = Workbook(); ws = wb.active; ws.title = "Transmittal"
    thin = Side(style="thin", color="000000")
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)
    black = PatternFill("solid", fgColor="000000")
    def F(sz=9, b=False, color="000000"): return Font(name=FONT, size=sz, bold=b, color=color)
    cen  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cenv = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left",   vertical="center")
    right= Alignment(horizontal="right",  vertical="center")

    NISS = len(issues); SPARE = max(2, 12 - NISS)   # keep ~12 issue columns like the template
    FIRST = 4; LAST = FIRST + NISS - 1; LASTALL = FIRST + NISS + SPARE - 1
    col = lambda i: get_column_letter(i)

    # ---- header block (matches spA template) ----
    ws.merge_cells("B1:C1"); ws["B1"] = "spaceagency"; ws["B1"].font = F(16); ws["B1"].alignment = right
    ws["A2"] = "DOCUMENT TRANSMITTAL"; ws["A2"].font = F(12); ws["A2"].alignment = left
    ws["A3"] = "To:";   ws["A3"].font = F(9, True);  ws.merge_cells("B3:C3")
    ws["A4"] = "Attn:"; ws["A4"].font = F(9, True);  ws.merge_cells("B4:C4")
    ws["A6"] = "Project Title";  ws["A6"].font = F(9, True);  ws.merge_cells("B6:C6"); ws["B6"] = title;  ws["B6"].font = F(9, True)
    ws["A7"] = "Project Number"; ws["A7"].font = F(9, True);  ws.merge_cells("B7:C7"); ws["B7"] = pno;    ws["B7"].font = F(9)
    ws.merge_cells(f"D7:{col(LASTALL)}7"); ws["D7"] = "Date Issued"; ws["D7"].font = F(9, True); ws["D7"].alignment = cenv
    ws.row_dimensions[1].height = 22

    # ---- column header row (row 8) ----
    HR = 8
    for lab, cix in [("DRAWING REGISTER", 1), ("No.", 2), ("CURRENT\nREV", 3)]:
        c = ws.cell(HR, cix, lab); c.font = F(9, True); c.alignment = cen; c.border = grid
    for k, iss in enumerate(issues):
        d = iss["date"]; c = ws.cell(HR, FIRST + k, f"{d[8:10]}.{d[5:7]}.{d[2:4]}")
        c.font = F(8, True); c.alignment = cen; c.border = grid
    for k in range(SPARE): ws.cell(HR, LAST + 1 + k).border = grid

    # ---- matrix body ----
    r = HR + 1
    def band(txt):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=LASTALL)
        c = ws.cell(r, 1, txt); c.font = F(9, True, "FFFFFF"); c.fill = black; c.alignment = left
        for ci in range(1, LASTALL + 1): ws.cell(r, ci).border = grid
        r += 1
    for series, nos in groups.items():
        band(series.upper())
        for no in nos:
            ws.cell(r, 1, cur[no]["name"]).font = F(9); ws.cell(r, 1).alignment = left; ws.cell(r, 1).border = grid
            ws.cell(r, 2, no).font = F(9, True); ws.cell(r, 2).alignment = cenv; ws.cell(r, 2).border = grid
            ws.cell(r, 3, cur[no]["rev"]).font = F(9, True); ws.cell(r, 3).alignment = cenv; ws.cell(r, 3).border = grid
            for k, iss in enumerate(issues):
                c = ws.cell(r, FIRST + k, iss["drawings"].get(no, "")); c.font = F(8); c.alignment = cenv; c.border = grid
            for k in range(SPARE): ws.cell(r, LAST + 1 + k).border = grid
            r += 1

    # ---- bottom sections ----
    r += 1
    def section_band(t):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=LASTALL)
        c = ws.cell(r, 1, t); c.font = F(9, True, "FFFFFF"); c.fill = black; c.alignment = left; r += 1
    section_band("Distribution")
    for _ in range(3):
        for ci in range(1, LASTALL + 1): ws.cell(r, ci).border = grid
        r += 1
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.cell(r, 1, "Method of Issue:  C - CD-Rom;  E - Email;  P - Post/Courier;  U - Upload;  H - Hardcopy").font = F(8, True)
    for k in range(NISS + SPARE): ws.cell(r, FIRST + k).border = grid
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.cell(r, 1, "Purpose of Issue:  A - Approval;  B - Building Permit;  C - Construction;  F - Forward Works;  I - Information;  R - Review;  T - For Tender;  P - Preliminary").font = F(8, True)
    for k in range(NISS + SPARE): ws.cell(r, FIRST + k).border = grid
    r += 2
    ws.cell(r, 1, "Issued by:").font = F(9, True); ws.cell(r, 2, "____________________").font = F(9)

    # ---- widths / freeze / print ----
    ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 10; ws.column_dimensions["C"].width = 8
    for k in range(NISS + SPARE): ws.column_dimensions[col(FIRST + k)].width = 6.5
    ws.freeze_panes = f"{col(FIRST)}{HR+1}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"1:{HR}"

    # ---- Notes / flags tab ----
    n = wb.create_sheet("Notes"); nr = 1
    def put(t, sz=9, b=False):
        nonlocal nr; c = n.cell(nr, 1, t); c.font = F(sz, b); c.alignment = left; nr += 1
    put(f"{title} — Drawing Issue Register / Transmittal", 12, True); nr += 1
    put("Auto-built from the project issue folders (build-from-folder mode).", 9)
    put("  • Rows = current drawings in __CURRENT DOCUMENTS/" + D["discipline"] + " (DELETED sheets excluded).", 9)
    put("  • Columns = each dated issue folder — the revision issued at that date, from the filename.", 9)
    nr += 1; put("Fill manually before issuing:", 10, True)
    put("  To / Attn / Distribution / Method / Purpose / Issued by; add the spA logo at top-left.", 9)
    put("  Revision notes (per rev) come from the titleblock revision history — added in a later pass.", 9)
    nr += 1; put("Flags — filenames missing a revision ID (correct at source):", 10, True)
    for mf in D["missing_rev"]: put("   ⚠  " + mf, 9)
    if issue_only:
        nr += 1; put("Issued but not in current set (check if superseded/renamed):", 10, True)
        for io in issue_only: put("   • " + io, 9)
    n.column_dimensions["A"].width = 115

    wb.save(a.out)
    print(f"saved {a.out}  ({len(rows)} drawings, {NISS} issues, {len(groups)} series, {len(D['missing_rev'])} missing-rev)")

if __name__ == "__main__":
    main()
